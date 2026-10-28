"""Build the ELECTRO-SOLVE 2026 deck: "warehouse ledger" light theme with Morph transitions.

Images come from Blender (build/blender/scene.py -> prep_assets.py) and the live app (build/shots).
Shapes whose names start with "!!" appear on every slide, so Morph glides them instead of re-drawing.
    python build/prep_assets.py && python build/build_deck.py
"""
import json
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "img"
import sys
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parent / "ElectroSolve_2026_Warehouse_Digital_Twin.pptx"
ROUTE = json.loads((HERE / "route_stats.json").read_text())

# ---------- design tokens (same as web/src/style.css) ----------
PAPER, SHEET = "F1EADB", "FAF7F0"
INK, INK2, MUTED, RULE = "1F2226", "4A4D52", "7D786C", "D3CAB8"
SIGNAL, COBALT, MOSS, PLUM, SAFETY = "E0442A", "2340B8", "2E7D4F", "A8307E", "E9AE0C"
F, F_LIGHT, F_SEMI, F_COND, MONO = "Bahnschrift", "Bahnschrift Light", "Bahnschrift SemiBold", "Bahnschrift SemiCondensed", "Cascadia Mono"
W, H = 13.333, 7.5
M = 0.6          # page margin
N_SLIDES = 11

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, name=None, kind=MSO_SHAPE.RECTANGLE, rot=0):
    sp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sp.name = name
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(lw)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    sp.rotation = rot
    return sp


def hline(s, x, y, w, color=INK, lw=0.75, name=None):
    return rect(s, x, y, w, lw / 72, fill=color, name=name)


def text(s, x, y, w, h, paras, size=14, color=INK, font=F, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spc=None, ls=None, after=0, name=None, rot=0):
    """paras: str | list[str | list[(text, {overrides})]]"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tb.rotation = rot
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p in enumerate([paras] if isinstance(paras, str) else paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if ls:
            para.line_spacing = ls
        para.space_after = Pt(after)
        for t, o in ([(p, {})] if isinstance(p, str) else p):
            r = para.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = rgb(o.get("color", color))
            if o.get("spc", spc):
                r._r.get_or_add_rPr().set("spc", str(o.get("spc", spc)))
    return tb


def pic(s, path, x, y, w=None, h=None, name=None):
    p = s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w) if w else None, height=Inches(h) if h else None)
    if name:
        p.name = name
    return p


def framed(s, path, x, y, w, name=None):
    """Screenshot with an ink hairline and the same hard offset shadow as the app's panels."""
    from PIL import Image
    iw, ih = Image.open(path).size
    h = w * ih / iw
    rect(s, x + 0.07, y + 0.07, w, h, fill=INK)
    p = pic(s, path, x, y, w, name=name)
    p.line.color.rgb = rgb(INK)
    p.line.width = Pt(1)
    return h


def add_morph(slide, dur_ms=1100):
    xml = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
           'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
           '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" '
           'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" Requires="p159">'
           f'<p:transition spd="slow" p14:dur="{dur_ms}"><p159:morph option="byObject"/></p:transition></mc:Choice>'
           '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
    el = etree.fromstring(xml)
    sld = slide._element
    anchor = sld.find(qn("p:clrMapOvr"))
    (anchor if anchor is not None else sld.find(qn("p:cSld"))).addnext(el)


# ---------- page furniture ----------
def page(idx, kicker=None, notes=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(PAPER)
    text(s, M, 0.34, 7, 0.25, "PICKING FLOOR DIGITAL TWIN", size=8.5, color=INK2, font=MONO, spc=120, name="!!run")
    text(s, W - M - 3, 0.34, 3, 0.25, [[(f"{idx + 1:02d}", {"color": SIGNAL}), (f" / {N_SLIDES:02d}", {})]],
         size=8.5, color=INK2, font=MONO, align=PP_ALIGN.RIGHT, name="!!pno")
    hline(s, M, 0.62, W - 2 * M, name="!!rule")
    hline(s, M, H - 0.52, W - 2 * M, color=RULE, name="!!foot")
    text(s, M, H - 0.42, 8, 0.22, "ELECTRO-SOLVE 2026  ·  PS 4.2  ·  SENSE, VIT", size=7.5, color=MUTED, font=MONO, spc=80, name="!!foottxt")
    text(s, W - M - 5, H - 0.42, 5, 0.22, "DATA: ONLINE RETAIL II (UCI, CC BY 4.0)", size=7.5, color=MUTED, font=MONO,
         spc=80, align=PP_ALIGN.RIGHT, name="!!src")
    # progress tick under the running header
    rect(s, M + (W - 2 * M) * idx / N_SLIDES, 0.6, (W - 2 * M) / N_SLIDES, 0.05, fill=SIGNAL, name="!!prog")
    if kicker:
        text(s, M, 1.0, 8, 0.3, [[(f"§{idx + 1:02d}  ", {"font": MONO, "color": MUTED, "spc": 0}), (kicker, {})]],
             size=11, color=SIGNAL, font=F_COND, bold=True, spc=220, name="!!kicker")
    s.notes_slide.notes_text_frame.text = notes
    if idx:
        add_morph(s)
    return s


def title(s, t, y=1.38, w=7.2, size=38, h=1.4):
    return text(s, M, y, w, h, t, size=size, font=F_SEMI, ls=0.92, name="!!title")


def body(s, t, x=M, y=2.75, w=5.4, size=14.5, color=INK2, h=1.4):
    return text(s, x, y, w, h, t, size=size, color=color, ls=1.12, name="!!body")


def big(s, x, y, w, value, label, color=SIGNAL, size=88, lsize=12.5, name=None):
    text(s, x, y, w, size / 72 + 0.1, value, size=size, color=color, font=F_LIGHT, name=name)
    text(s, x, y + size / 72 + 0.12, w, 0.6, label, size=lsize, color=INK2, ls=1.1)


def ledger(s, x, y, w, rows, active=None, scale=None, unit="%", label_w=1.35, val_w=0.95, row_h=0.46):
    """Results ledger: label | bar | value, hairlines between rows."""
    vmax = scale or max(v for _, v, _ in rows)
    hline(s, x, y, w)
    for i, (lab, v, vt) in enumerate(rows):
        yy = y + i * row_h
        on = active is None or i == active
        text(s, x, yy + 0.1, label_w, 0.3, lab, size=12.5, font=F_SEMI if i == active else F, color=INK if on else MUTED)
        bw = (w - label_w - val_w - 0.15) * (v / vmax if vmax else 0)
        rect(s, x + label_w, yy + 0.17, max(bw, 0.03), 0.13, fill=SIGNAL if i == active else (INK if on else "BDB4A3"))
        text(s, x + w - val_w, yy + 0.08, val_w, 0.32, vt, size=13, font=MONO, color=SIGNAL if i == active else (INK if on else MUTED),
             align=PP_ALIGN.RIGHT)
        hline(s, x, yy + row_h, w, color=RULE)


LADDER = [("Original", 0, "218.1 km"), ("ABC", 2.5, "−2.5%"), ("Affinity", 11.3, "−11.3%"), ("Optimized", 15.8, "−15.8%")]

NOTES = [
    "Hook: pickers in a warehouse spend more time walking than picking, and where each product sits decides how far they walk. "
    "We built a 3D digital twin of a picking floor, driven by 2,293 real customer orders. Rearrange the shelves and the twin replays "
    "exactly the same orders to measure the walking saved, instead of just claiming it.",
    "In our model one order costs about 79 s of walking and 71 s of picking, so 53% of the time is walking. Over the month one picker "
    f"walks 218 km in a 100-slot zone. The picture is one real order, invoice {ROUTE['order']}: {ROUTE['lines']} lines, and on the "
    f"as-received layout the picker walks {ROUTE['original']['m']} m through all {ROUTE['original']['aisles']} aisles.",
    "Dataset: UCI Online Retail II, a UK gift wholesaler, November 2010, the pre-Christmas peak. We remove cancellations, postage and fees, "
    "and invalid rows: 76,584 lines become 74,757. We keep the top 100 products by pick lines: 3% of the range but 22% of picks. "
    "Result: 2,293 real orders, 7.1 lines per order. A pick is one order line, because walking depends on visits, not quantity.",
    "Architecture: Python cleans the data once; one shared JavaScript simulation core routes and scores any layout. The same code runs "
    "offline in Node for the optimizer and live in the browser for the game, so the numbers are identical. The floor is a graph of "
    "61 nodes and 64 edges, and no edge passes through a rack. The 3D floor itself is baked in Blender from the same geometry.",
    "The textbook fix is ABC slotting: most-picked products nearest the dock. On the same 2,293 orders it saves only 2.5%. Why: an "
    "average order has 7 lines spread over many products, so the picker still enters most aisles. Rack colour is how often each slot is visited.",
    "Affinity slotting uses what customers buy together: products that share orders are pulled into the same aisle, while fast movers "
    "still sit near the dock. Same orders: 218.1 to 193.4 km, 11.3% less walking, four to five times the ABC saving.",
    "Our optimizer is simulated annealing over slot swaps: 80,000 swaps, each judged by re-routing every affected order of the month. "
    "462 were accepted, in about 44 seconds on a laptop. Result: 183.7 km, 15.8% less walking. That is 34 km, about 8 hours of picker time a month.",
    f"Here is what that means for one order. Invoice {ROUTE['order']}, {ROUTE['lines']} lines. Original layout: {ROUTE['original']['m']} m "
    f"through {ROUTE['original']['aisles']} aisles. Optimized: {ROUTE['optimized']['m']} m through {ROUTE['optimized']['aisles']}. "
    "To be fair, this is one of the biggest single-order wins of the month; on average an order goes from 95 to 80 m.",
    "Is it memorising the month? No. We built the layouts only from orders before 22 November and tested on the final week, 766 orders "
    "they never saw: ABC 1.9%, Affinity 11.2%, Optimized 13.9%. And slotting stacks with process: batching 4 orders per trip cuts the "
    "month to 88 km; layout-aware smart batching plus the optimized layout gets to 70 km, 68% less than where we started.",
    "Demo. Click one rack, then another, to swap; the whole month re-scores instantly. Diamonds show products bought together. "
    "The live optimizer runs the same annealing in a background thread on your layout; in this screenshot it beat the offline optimizer, "
    "182.3 km. Space to replay, 1 to 5 layouts, H heatmap, B batching, I IoT mode, O optimizer.",
    "ECE angle: a digital twin needs sensors. The replay already emits the events real devices would stream: an RFID gate at the dock, "
    "a UWB tag on the cart, pick-to-light modules on every slot, and barcode confirmations. It downloads as JSON Lines, so real hardware "
    "can replace the replay. Next: several pickers and aisle congestion. Thank you, questions welcome.",
]

# =========================== slides ===========================
# 1 — title
s = page(0, notes=NOTES[0])
pic(s, IMG / "axo_optimized.jpg", 5.0, 0.75, w=8.9, name="!!hero")
rect(s, M, 1.05, 0.95, 0.36, line=SIGNAL, lw=1.5, rot=-3)
text(s, M, 1.05, 0.95, 0.36, "PS 4.2", size=11, color=SIGNAL, font=F_COND, bold=True, align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE, spc=150, rot=-3)
text(s, M + 1.15, 1.13, 4, 0.3, "ELECTRO-SOLVE 2026  ·  PROJECT EXPO", size=10, color=INK2, font=MONO, spc=100)
text(s, M, 1.75, 6, 2.3, [[("Picking Floor", {"font": F_SEMI})], [("Digital Twin", {"font": F_LIGHT, "color": SIGNAL})]],
     size=60, ls=0.9, name="!!title")
body(s, "A 3D twin of a warehouse picking floor, driven by 2,293 real customer orders. Re-shelve the products and it replays "
        "the same month to measure the walking saved.", y=3.85, w=4.6, size=15)
for i, (v, l) in enumerate([("2,293", "real orders replayed"), ("−15.8%", "walking, optimized layout"), ("−68%", "with smart batching")]):
    x = M + i * 1.62
    hline(s, x, 5.35, 1.45)
    text(s, x, 5.45, 1.5, 0.5, v, size=24, font=F_LIGHT, color=SIGNAL if i else INK)
    text(s, x, 5.95, 1.45, 0.5, l, size=9.5, color=INK2, ls=1.05)
text(s, M, 6.3, 6.2, 0.7, [[("Team GitHappens  ·  SENSE, VIT", {"font": F_SEMI, "color": INK})],
                           [("Rohith Krishna A K (25BCE2027)  ·  Jaidev A (25BCE0971)", {})],
                           [("Guide: Dr. Jaffino", {})]], size=10.5, color=MUTED)

# 2 — problem
s = page(1, "THE PROBLEM", NOTES[1])
title(s, "Pickers walk more\nthan they pick.")
pic(s, IMG / "route_original.jpg", 5.95, 0.85, w=7.2, name="!!hero")
big(s, M, 2.85, 3.0, "53%", "of every order’s time is walking", size=96)
# time split bar: 79 s walking / 71 s picking
x0, wbar = M, 4.9
rect(s, x0, 4.95, wbar * 79 / 150, 0.42, fill=INK)
rect(s, x0 + wbar * 79 / 150, 4.95, wbar * 71 / 150, 0.42, line=INK, lw=1)
text(s, x0 + 0.12, 4.95, 2.4, 0.42, "79 s walking", size=12, color=PAPER, font=F_SEMI, anchor=MSO_ANCHOR.MIDDLE)
text(s, x0 + wbar * 79 / 150 + 0.12, 4.95, 2.2, 0.42, "71 s picking", size=12, font=F_SEMI, anchor=MSO_ANCHOR.MIDDLE)
text(s, M, 5.48, 5, 0.3, "Per order, 1.2 m/s walking, 10 s per line", size=10, color=MUTED, font=MONO)
text(s, M, 5.95, 5.0, 0.8, [[("218 km", {"font": F_SEMI, "color": INK}), (" walked in one month by a single picker in a 100-slot zone.", {})]],
     size=13, color=INK2)
text(s, 6.35, 6.35, 6.4, 0.45,
     [[(f"Invoice {ROUTE['order']} · {ROUTE['lines']} lines", {"font": F_SEMI, "color": INK}),
       (f"  on the as-received layout: {ROUTE['original']['m']} m through all {ROUTE['original']['aisles']} aisles", {})]],
     size=10.5, color=INK2, align=PP_ALIGN.RIGHT)

# 3 — data
s = page(2, "THE DATA", NOTES[2])
title(s, "Built on 2,293 real customer orders", w=11)
body(s, "Online Retail II (UCI, CC BY 4.0): a UK gift wholesaler, November 2010, the pre-Christmas peak.", y=2.1, w=9)
funnel = [("Raw order lines", 76584, "76,584"), ("Clean pick lines", 74757, "74,757"),
          ("Lines on the top 100 products", 16225, "16,225"), ("Real orders replayed", 2293, "2,293")]
fy = 3.05
hline(s, M, fy, 7.4)
for i, (lab, v, vt) in enumerate(funnel):
    y = fy + i * 0.78
    text(s, M, y + 0.14, 2.6, 0.5, lab, size=12.5, color=INK if i == 3 else INK2, font=F_SEMI if i == 3 else F)
    rect(s, M + 2.75, y + 0.2, max(3.45 * v / 76584, 0.04), 0.36, fill=SIGNAL if i == 3 else (INK if i < 2 else INK2))
    text(s, M + 6.3, y + 0.17, 1.1, 0.4, vt, size=13.5, font=MONO, color=SIGNAL if i == 3 else INK, align=PP_ALIGN.RIGHT)
    hline(s, M, y + 0.78, 7.4, color=RULE)
notes3 = ["−1,192 cancelled", "−282 postage & fees · −353 invalid", "top 100 of 3,018 products", "7.1 lines per order"]
for i, n in enumerate(notes3):
    text(s, M + 2.75, fy + i * 0.78 + 0.58, 4, 0.2, n, size=8, color=MUTED, font=MONO)
big(s, 8.75, 2.95, 4.0, "2,293", "real orders, replayed line by line", size=80, lsize=14)
text(s, 8.75, 4.75, 3.9, 1.6, [[("3%", {"font": F_SEMI, "color": INK}), (" of the product range carries ", {}),
                               ("22%", {"font": F_SEMI, "color": INK}), (" of all picks. A pick is one order line: walking depends on visits, not on quantity.", {})]],
     size=13, color=INK2, ls=1.15)

# 4 — how it works
s = page(3, "HOW IT WORKS", NOTES[3])
title(s, "One simulation core, used twice", w=11)
body(s, "The optimizer and the live game run the same routing code, so every number on screen is measured, not claimed.", y=2.1, w=9)
steps = [("Clean", "pandas turns 76,584 raw lines into 2,293 orders on the top 100 products."),
         ("Floor graph", "61 nodes, 64 edges. No edge crosses a rack: walking follows the aisles."),
         ("Route & score", "Shortest paths + nearest-neighbour and 2-opt per trip. A month replays in milliseconds."),
         ("Optimize & play", "Simulated annealing offline and in a Web Worker; Three.js twin in the browser.")]
cw = (W - 2 * M) / 4
for i, (h_, d) in enumerate(steps):
    x = M + i * cw
    hline(s, x, 3.05, cw - 0.25, color=INK)
    text(s, x, 3.15, 1.5, 0.9, f"{i + 1:02d}", size=48, font=F_LIGHT, color=SIGNAL)
    text(s, x, 4.1, cw - 0.4, 0.4, h_, size=17, font=F_SEMI)
    text(s, x, 4.55, cw - 0.45, 1.3, d, size=12.5, color=INK2, ls=1.12)
rect(s, M, 5.85, W - 2 * M, 0.62, fill=INK)
text(s, M + 0.25, 5.85, W - 2 * M - 0.5, 0.62,
     [[("Same deterministic code in Node and in the browser  →  ", {"color": "F3B8A9", "font": F_SEMI}),
       ("both agree to the metre: 218.1 km for Original.", {})]],
     size=13, color=PAPER, anchor=MSO_ANCHOR.MIDDLE)

# 5–7 — the layout ladder
LAD = [
    ("LAYOUT 1  ·  INTUITION", "The textbook fix\nsaves only 2.5%", "axo_abc.jpg", 1,
     "ABC slotting puts the most-picked products nearest the dock. Orders average 7 lines, so the picker still visits most aisles.",
     "−2.5%", "218.1 → 212.7 km a month"),
    ("LAYOUT 2  ·  DATA INSIGHT", "What customers buy\ntogether wins", "axo_affinity.jpg", 2,
     "Affinity slotting pulls co-ordered products into the same aisle, while fast movers still sit near the dock.",
     "−11.3%", "218.1 → 193.4 km · 4–5× the ABC saving"),
    ("LAYOUT 3  ·  ALGORITHM", "80,000 swaps later:\n15.8% less walking", "axo_optimized.jpg", 3,
     "Simulated annealing tries random slot swaps; each is judged by re-routing every affected order of the month.",
     "−15.8%", "218.1 → 183.7 km · 34 km ≈ 8 h of picker time"),
]
for j, (kick, ttl, img, act, bd, num, sub) in enumerate(LAD):
    s = page(4 + j, kick, NOTES[4 + j])
    pic(s, IMG / img, 5.35, 0.8, w=8.0, name="!!hero")
    title(s, ttl, w=5.3, size=34)
    body(s, bd, y=2.6, w=4.7, size=13.5)
    text(s, M, 3.75, 3, 0.9, num, size=60, font=F_LIGHT, color=SIGNAL, name="!!num")
    text(s, M, 4.68, 4.6, 0.3, sub, size=10.5, color=INK2, font=MONO, name="!!sub")
    ledger(s, M, 5.12, 4.65, LADDER, active=act, scale=15.8, row_h=0.36)
    text(s, 7.2, 6.55, 5.55, 0.3, "Rack colour = visits per month: rare → very frequent", size=8.5, color=MUTED,
         font=MONO, align=PP_ALIGN.RIGHT)
    for k in range(5):  # tiny ramp key
        rect(s, 10.65 + k * 0.42, 6.38, 0.42, 0.1, fill=["D7D3C8", "A9C3BB", "4E978A", "EDB23A", "E0442A"][k])

# 8 — same order, two floors
s = page(7, "ONE REAL ORDER", NOTES[7])
title(s, "Same order, two floors", w=11)
body(s, f"Invoice {ROUTE['order']}: {ROUTE['lines']} lines, the same picker, routed along the aisles on each layout.", y=2.05, w=9)
for i, (k, lab) in enumerate((("original", "ORIGINAL"), ("optimized", "OPTIMIZED"))):
    x = M + i * 6.17
    pic(s, IMG / f"route_{k}.jpg", x - 0.1, 2.55, w=6.2, name="!!hero" if i else None)
    r = ROUTE[k]
    hline(s, x, 6.0, 5.9, color=INK)
    text(s, x, 6.08, 2, 0.3, lab, size=10.5, font=F_COND, bold=True, spc=200, color=SIGNAL if i else INK2)
    text(s, x + 2.0, 6.0, 3.9, 0.5, [[(f"{r['m']} m", {"font": F_LIGHT, "size": 26, "color": SIGNAL if i else INK}),
                                      (f"   {r['aisles']} aisle{'s' if r['aisles'] > 1 else ''}", {"color": INK2})]],
         size=13, align=PP_ALIGN.RIGHT)
text(s, M, 6.55, 12, 0.3, "One of the month’s biggest single-order wins. Averaged over all 2,293 orders: 95.1 → 80.1 m per order.",
     size=9.5, color=MUTED, font=MONO)

# 9 — proof + batching
s = page(8, "PROOF  ·  UNSEEN ORDERS", NOTES[8])
title(s, "It holds on orders it never saw,\nand stacks with batching", w=11, size=34)
# left: hold-out grouped bars
lx, ly, lw_ = M, 2.95, 5.6
text(s, lx, ly, lw_, 0.3, "WALKING SAVED VS ORIGINAL", size=10, font=F_COND, bold=True, spc=180, color=INK2)
rect(s, lx + 3.35, ly + 0.05, 0.16, 0.16, fill=INK)
text(s, lx + 3.56, ly, 1, 0.3, "full month", size=9.5, color=INK2)
rect(s, lx + 4.55, ly + 0.05, 0.16, 0.16, fill=COBALT)
text(s, lx + 4.76, ly, 1, 0.3, "unseen week", size=9.5, color=INK2)
hline(s, lx, ly + 0.42, lw_)
for i, (lab, a, b) in enumerate([("ABC", 2.5, 1.9), ("Affinity", 11.3, 11.2), ("Optimized", 15.8, 13.9)]):
    y = ly + 0.55 + i * 0.85
    text(s, lx, y + 0.12, 1.3, 0.4, lab, size=13, font=F_SEMI if i == 2 else F)
    for jj, (v, c) in enumerate(((a, INK), (b, COBALT))):
        bw = 3.4 * v / 15.8
        rect(s, lx + 1.35, y + jj * 0.3, bw, 0.22, fill=c)
        text(s, lx + 1.45 + bw, y + jj * 0.3 - 0.04, 0.9, 0.3, f"{v:.1f}%", size=10.5, font=MONO, color=c)
    hline(s, lx, y + 0.72, lw_, color=RULE)
text(s, lx, 5.95, lw_, 0.5, "Layouts built on orders before 22 Nov (1,527), tested on the final week (766 orders).",
     size=10, color=MUTED, ls=1.1)
# right: the walking staircase
rx, rw = 7.1, 5.63
text(s, rx, ly, rw, 0.3, "KM WALKED IN THE MONTH", size=10, font=F_COND, bold=True, spc=180, color=INK2)
hline(s, rx, ly + 0.42, rw)
stair = [("Original, 1 order / trip", 218.1, INK2), ("Optimized layout", 183.7, INK), ("Batch of 4 orders", 88.2, INK),
         ("Smart batch + optimized", 70.1, SIGNAL)]
for i, (lab, v, c) in enumerate(stair):
    y = ly + 0.52 + i * 0.52
    text(s, rx, y + 0.06, 2.3, 0.4, lab, size=11.5, color=INK if i == 3 else INK2, font=F_SEMI if i == 3 else F)
    rect(s, rx + 2.35, y + 0.1, 2.5 * v / 218.1, 0.26, fill=c)
    text(s, rx + 2.45 + 2.5 * v / 218.1, y + 0.05, 1, 0.32, f"{v:.1f}", size=11.5, font=MONO, color=c)
text(s, rx, 5.25, 3, 1.0, "−68%", size=54, font=F_LIGHT, color=SIGNAL)
text(s, rx + 2.25, 5.42, 3.3, 0.9, "walking: the best layout plus layout-aware smart batching, on the same orders.",
     size=11.5, color=INK2, ls=1.1)

# 10 — play it
s = page(9, "PLAY IT", NOTES[9])
title(s, "Can you beat\nthe optimizer?", w=4.6, size=36)
framed(s, IMG / "app_game.jpg", 5.25, 1.0, 7.45, name="!!hero")
feats = [("Swap racks", "The month re-scores instantly; ◆ marks what is bought together."),
         ("Live optimizer", "Annealing in a Web Worker on your layout. Here: 182.3 km in 14 s."),
         ("See the traffic", "Rack visit colours and a floor heatmap of walking passes."),
         ("IoT sensor mode", "RFID, UWB, pick-to-light and scan events, exported as JSON Lines.")]
for i, (a, b) in enumerate(feats):
    y = 2.75 + i * 0.86
    hline(s, M, y, 4.3, color=INK if i == 0 else RULE)
    text(s, M, y + 0.1, 0.5, 0.3, f"{i + 1:02d}", size=10, font=MONO, color=SIGNAL)
    text(s, M + 0.5, y + 0.07, 3.8, 0.3, a, size=14, font=F_SEMI)
    text(s, M + 0.5, y + 0.4, 3.8, 0.45, b, size=10.5, color=INK2, ls=1.05)
text(s, M, 6.3, 4.4, 0.3, "Space replay · 1–5 layouts · H heat · B batch · I IoT · O optimize", size=8.5, color=MUTED, font=MONO)

# 11 — close
s = page(10, "WHAT'S NEXT", NOTES[10])
title(s, "From replay\nto a live twin", w=5.2, size=38)
body(s, "A twin needs sensors. The replay already emits the events real devices would stream, so hardware can replace it one feed at a time.",
     y=2.75, w=4.8, size=13.5)
framed(s, IMG / "detail.jpg", 6.1, 1.0, 6.55, name="!!hero")
cards = [("RFID", "Dock gate", "cart_out C-01 · 4 totes"), ("UWB", "Cart tag", "x 12.5 z 1.0 · aisle 3"),
         ("PICK-TO-LIGHT", "Every slot", "light_on A3-R05 · qty 6"), ("BARCODE", "Picker scanner", "pick_confirmed 22694 ✓")]
cw = (W - 2 * M - 0.45) / 4
for i, (t, d, ev) in enumerate(cards):
    x = M + i * (cw + 0.15)
    rect(s, x, 5.75, cw, 0.85, fill=SHEET, line=INK, lw=0.75)
    text(s, x + 0.15, 5.83, cw - 0.3, 0.3, [[(t, {"color": [PLUM, COBALT, MOSS, SIGNAL][i], "font": F_COND, "bold": True}), (f"  {d}", {"color": INK2})]],
         size=11, spc=60)
    text(s, x + 0.15, 6.2, cw - 0.3, 0.3, ev, size=9.5, font=MONO, color=INK)
text(s, M, 4.25, 5.2, 0.7, "Thank you. Questions?", size=26, font=F_LIGHT, color=SIGNAL)
text(s, M, 4.95, 5.2, 0.4, "Next: several pickers and aisle congestion.", size=11.5, color=INK2)

prs.save(OUT)
print("saved", OUT, len(prs.slides._sldIdLst), "slides")
