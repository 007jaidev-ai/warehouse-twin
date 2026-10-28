"""Official-format A3 poster (SENSE template 25EUEC1001HW-Poster-Template.pptx), filled with this project.

Keeps the template's red frame and VIT logo, follows its section order and type sizes
(title Calibri 24, names 21, headers 18, body 12, captions 10). Tables, the chart and Fig. 1 are
native PowerPoint objects, so everything stays editable.
    python build/prep_assets.py && node web/scripts/congestion.mjs && python build/build_official_poster.py
"""
import copy
import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
ROOT = HERE.parent
TEMPLATE = ROOT / "25EUEC1001HW-Poster-Template.pptx"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "ElectroSolve_2026_Official_Poster_GitHappens.pptx"
IMG = HERE / "img"
CONG = json.loads((ROOT / "web" / "public" / "data" / "congestion.json").read_text())
ROUTE = json.loads((HERE / "route_stats.json").read_text())

NAVY, RULE, INK, MUTED = "1F3864", "42719B", "000000", "595959"
SIGNAL, COBALT, MOSS = "E0442A", "2340B8", "2E7D4F"
F = "Calibri"

prs = Presentation(TEMPLATE)
s = prs.slides[0]
KEEP = {"Google Shape;1122;p13", "Google Shape;1137;p13"}  # red frame, VIT logo
for sh in list(s.shapes):
    if sh.name not in KEEP:
        sh._element.getparent().remove(sh._element)


def rgb(h):
    return RGBColor.from_string(h)


def tb(x, y, w, h, paras, size=12, color=INK, bold=False, align=PP_ALIGN.JUSTIFY, border=False, anchor=MSO_ANCHOR.TOP,
       ls=0.92, inset=0.07, after=0):
    """paras: list of str | list[(text, {bold, color, size, italic})]. Returns the shape."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(inset)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    if border:
        box.line.color.rgb = rgb(RULE)
        box.line.width = Pt(1.25)
    for i, p in enumerate([paras] if isinstance(paras, str) else paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = ls
        para.space_after = Pt(after)
        for t, o in ([(p, {})] if isinstance(p, str) else p):
            r = para.add_run()
            r.text = t
            r.font.name = F
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = rgb(o.get("color", color))
    return box


def header(x, y, text, w=5.4):
    return tb(x, y, w, 0.36, text, size=18, bold=True, color=NAVY, align=PP_ALIGN.LEFT, inset=0.02)


def caption(x, y, w, text):
    return tb(x, y, w, 0.25, text, size=10, color=INK, align=PP_ALIGN.CENTER, inset=0.02)


def picture(path, x, y, w=None, h=None, border=True):
    p = s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w) if w else None, height=Inches(h) if h else None)
    if border:
        p.line.color.rgb = rgb("BFBFBF")
        p.line.width = Pt(0.75)
    return p


def box(x, y, w, h, fill, line=None, text=None, size=9, color="FFFFFF", bold=True, kind=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = rgb(fill)
    if line:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(1)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = 0.12
    if text:
        tf = sp.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line_ in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.line_spacing = 0.9
            r = p.add_run()
            r.text = line_
            r.font.name = F
            r.font.size = Pt(size if i == 0 else size - 1)
            r.font.bold = bold if i == 0 else False
            r.font.color.rgb = rgb(color)
    return sp


def arrow(x1, y1, x2, y2, color=NAVY):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


# ================= title block (template positions) =================
tb(2.83, 0.25, 8.55, 0.5, [[("Picking Floor Digital Twin: Data-Driven Warehouse Slotting", {})]],
   size=24, bold=True, align=PP_ALIGN.LEFT, inset=0.05)
tb(2.83, 0.72, 8.55, 0.75, [[("Rohith Krishna A K (25BCE2027) and Jaidev A (25BCE0971)  |  Team GitHappens", {"size": 18})],
                            [("Dr. Jaffino  |  SENSE", {})]], size=21, align=PP_ALIGN.LEFT, inset=0.05)

L, R, CW = 0.33, 5.95, 5.38   # column x positions and width

# ================= left column =================
header(L - 0.05, 1.42, "Motivation/ Introduction")
tb(L, 1.78, CW, 1.1, [[
    ("In a warehouse, an order picker spends more time ", {}), ("walking", {"bold": True}),
    (" than picking: in our model 79 s of every 150 s order is walking. Where each product sits on the shelves (", {}),
    ("slotting", {"bold": True}),
    (") decides that distance, yet it is usually set by habit. We built a ", {}), ("digital twin", {"bold": True}),
    (" of a picking floor that replays real customer orders, so any shelf arrangement can be measured before anyone moves a carton.", {})]],
   border=True)

header(L - 0.05, 2.95, "SCOPE of the Project")
tb(L, 3.31, CW, 1.55, [[
    ("A 100-slot pick zone (5 aisles, racks on both sides) is driven by ", {}), ("2,293 real orders", {"bold": True}),
    (" from the UCI Online Retail II dataset (Nov 2010). We compare four slotting strategies, batch and smart-batch picking, and a ", {}),
    ("team of up to six pickers with aisle congestion", {"bold": True}),
    (". The twin runs in a browser as a 3D game where users swap racks against an optimizer, and it emits the events real sensors "
     "would stream: RFID dock gate, UWB cart tags, pick-to-light modules and barcode scans.", {})]], border=True)

header(L - 0.05, 4.93, "Methodology")
# ---- Fig. 1: pipeline (two rows of four, native shapes)
bw, bh, gap = 1.19, 0.72, 0.2
row1 = [("Real orders", "UCI Online Retail II\n76,584 lines"), ("Clean (pandas)", "2,293 orders\ntop 100 products"),
        ("Floor graph", "61 nodes, 64 edges\nFloyd–Warshall"), ("Route each trip", "nearest neighbour\n+ 2-opt")]
row2 = [("IoT event stream", "RFID · UWB · PTL\nbarcode, JSONL"), ("3D twin + game", "Three.js, Web Worker\nBlender-baked floor"),
        ("Team of pickers", "aisle capacity\nqueueing sim"), ("Slotting", "ABC · Affinity\nsimulated annealing")]
fy1, fy2 = 5.36, 6.42
fills1 = ["7F7F7F", NAVY, NAVY, NAVY]
fills2 = [MOSS, COBALT, SIGNAL, SIGNAL]
for i, (a, b) in enumerate(row1):
    x = L + 0.05 + i * (bw + gap)
    box(x, fy1, bw, bh, fills1[i], text=f"{a}\n{b}")
    if i < 3:
        arrow(x + bw + 0.02, fy1 + bh / 2, x + bw + gap - 0.02, fy1 + bh / 2)
for i, (a, b) in enumerate(row2):
    x = L + 0.05 + i * (bw + gap)
    box(x, fy2, bw, bh, fills2[i], text=f"{a}\n{b}")
    if i < 3:
        arrow(x + bw + gap - 0.02, fy2 + bh / 2, x + bw + 0.02, fy2 + bh / 2)
xr = L + 0.05 + 3 * (bw + gap) + bw / 2
arrow(xr, fy1 + bh + 0.02, xr, fy2 - 0.02)
caption(L, 7.17, CW, "Fig. 1  System pipeline: one simulation core, used offline (Node) and live (browser).")

tb(L, 7.47, CW, 5.22, [
    [("a)  Data. ", {"bold": True}), ("Cancelled invoices, postage/fee codes and invalid rows are removed (76,584 → 74,757 lines). "
     "The top 100 products by pick lines (3% of the range, 22% of picks) fill the 100 slots; 2,293 orders remain, 7.1 lines each.", {})],
    [("b)  Floor and routing. ", {"bold": True}), ("The floor is a graph whose edges follow aisles, so no route can cut through a rack. "
     "Floyd–Warshall gives all shortest paths; each trip (dock → picks → dock) is ordered by nearest neighbour and improved by 2-opt. "
     "Walking 1.2 m/s, 10 s per order line.", {})],
    [("c)  Slotting. ", {"bold": True}), ("Original (stock-code order); ABC (most-picked nearest the dock); Affinity (slot score = picks + "
     "w × co-orders with products already in that aisle); Optimized (simulated annealing over slot swaps, 80,000 iterations, "
     "accepting a worse swap with probability ", {}), ("P = e^(−Δd / T)", {"bold": True, "italic": True}),
     ("; only orders touching the two swapped products are re-routed).", {})],
    [("d)  Team of pickers. ", {"bold": True}), ("N pickers (1–6) clear the month's backlog from one queue. An aisle holds at most c carts "
     "(narrow: c = 1, wide: c = 2); a picker reaching a full aisle waits at its entrance. Time-stepped (0.5 s) and deterministic, "
     "so the browser and Node give identical results. Congestion loss = 1 − throughput / (N × single-picker throughput).", {})],
    [("e)  Validation. ", {"bold": True}), ("Layouts were also built from orders before 22 Nov only and tested on the 766 orders of the final week.", {})],
    [("f)  IoT layer. ", {"bold": True}), ("Every replay emits the events real devices would publish: RFID gate at the dock (cart out/in), "
     "UWB tag position at every aisle node, pick-to-light ON/OFF per slot, barcode pick confirmation and aisle-blocked alerts, "
     "exported as JSON Lines. An ESP32 publishing the same messages (e.g. over MQTT) could replace the replay feed.", {})],
], size=11.5, border=True, after=3)

picture(IMG / "app_team.jpg", L, 12.82, w=CW)
caption(L, 15.86, CW, "Fig. 5  3D twin in team mode: two of four pickers queue at a full aisle.")

# ================= right column =================
header(R - 0.05, 1.42, "Results")
iw = (CW - 0.1) / 2
picture(IMG / "route_original.jpg", R, 1.82, w=iw)
picture(IMG / "route_optimized.jpg", R + iw + 0.1, 1.82, w=iw)
caption(R, 1.82 + iw * 0.625 + 0.02, CW,
        f"Fig. 2  One real order (invoice {ROUTE['order']}, {ROUTE['lines']} lines): {ROUTE['original']['m']} m through "
        f"{ROUTE['original']['aisles']} aisles on Original (left) vs {ROUTE['optimized']['m']} m through {ROUTE['optimized']['aisles']} on Optimized.")

ty = 3.88
rows = [("Layout", "km / month", "m / order", "Walking saved", "Unseen week"),
        ("Original", "218.1", "95.1", "–", "–"),
        ("ABC (intuition)", "212.7", "92.8", "2.5%", "1.9%"),
        ("Affinity (data)", "193.4", "84.3", "11.3%", "11.2%"),
        ("Optimized (SA)", "183.7", "80.1", "15.8%", "13.9%"),
        ("Smart batch + Optimized", "70.1", "30.6", "67.9%", "–")]
tbl = s.shapes.add_table(len(rows), 5, Inches(R), Inches(ty), Inches(CW), Inches(0.24 * len(rows))).table
for j, wdt in enumerate([1.9, 0.85, 0.85, 0.95, 0.83]):
    tbl.columns[j].width = Inches(wdt)
for i, r in enumerate(rows):
    tbl.rows[i].height = Inches(0.24)
    for j, v in enumerate(r):
        c = tbl.cell(i, j)
        c.margin_left = c.margin_right = Inches(0.05)
        c.margin_top = c.margin_bottom = Inches(0.01)
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(NAVY if i == 0 else ("FBE3DE" if i == 4 else ("FFFFFF" if i % 2 else "EEF2F8")))
        p = c.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
        run = p.add_run()
        run.text = v
        run.font.name = F
        run.font.size = Pt(10)
        run.font.bold = i == 0 or i == 4
        run.font.color.rgb = rgb("FFFFFF" if i == 0 else INK)
caption(R, ty + 0.24 * len(rows) + 0.03, CW,
        "Table 1  Same 2,293 orders on every layout, one picker. Unseen week: layouts built before 22 Nov, tested on the last 766 orders.")

# ---- Fig. 3: team of pickers (native line chart)
cy = 5.75
opt1, opt2 = CONG["caps"]["1"]["optimized"], CONG["caps"]["2"]["optimized"]
one = opt1[0]["ordersPerHour"]
cd = CategoryChartData()
cd.categories = [str(n) for n in CONG["pickers"]]
cd.add_series("Perfect scaling", [round(one * n, 1) for n in CONG["pickers"]])
cd.add_series("Wide aisles (2 carts)", [round(r["ordersPerHour"], 1) for r in opt2])
cd.add_series("Narrow aisles (1 cart)", [round(r["ordersPerHour"], 1) for r in opt1])
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(R), Inches(cy), Inches(3.35), Inches(2.25), cd)
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(8)
ch.legend.font.name = F
ch.font.name = F
ch.font.size = Pt(9)
for ser, col, dash in zip(ch.plots[0].series, ["7F7F7F", COBALT, SIGNAL], [True, False, False]):
    ser.format.line.color.rgb = rgb(col)
    ser.format.line.width = Pt(2.25)
    ser.smooth = False
    if dash:
        ser.format.line.dash_style = 4  # dash
    ser.marker.style = XL_MARKER_STYLE.CIRCLE
    ser.marker.size = 5
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = rgb(col)
    ser.marker.format.line.color.rgb = rgb(col)
va = ch.value_axis
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = rgb("D9D9D9")
va.has_title = True
va.axis_title.text_frame.text = "orders / hour"
va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
ca = ch.category_axis
ca.has_title = True
ca.axis_title.text_frame.text = "pickers on the floor"
ca.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
n4, n6 = opt1[3], opt1[5]
loss4 = 1 - n4["ordersPerHour"] / (4 * one)
tb(R + 3.42, cy + 0.05, CW - 3.42, 2.2, [
    [("4 pickers, narrow aisles", {"bold": True, "size": 10.5})],
    [(f"{loss4 * 100:.0f}% of team capacity lost; pickers wait {n4['waitShare'] * 100:.0f}% of the time.", {"size": 10.5})],
    [("6 pickers", {"bold": True, "size": 10.5})],
    [(f"month cleared in {n6['hours']:.1f} h (narrow) vs {opt2[5]['hours']:.1f} h (wide aisles).", {"size": 10.5})],
    [("Slotting vs congestion", {"bold": True, "size": 10.5})],
    [(f"Optimized beats Original by +{(opt1[0]['ordersPerHour'] / CONG['caps']['1']['original'][0]['ordersPerHour'] - 1) * 100:.0f}% "
      f"with 1 picker, +{(n6['ordersPerHour'] / CONG['caps']['1']['original'][5]['ordersPerHour'] - 1) * 100:.0f}% with 6.", {"size": 10.5})],
], align=PP_ALIGN.LEFT, ls=0.95, after=2)
caption(R, cy + 2.27, CW, "Fig. 3  Throughput vs team size on the Optimized layout.")

# ---- Fig. 4: the four layouts as plans
py = 8.42
names = [("original", "Original 218.1 km"), ("abc", "ABC −2.5%"), ("affinity", "Affinity −11.3%"), ("optimized", "Optimized −15.8%")]
pw = (CW - 0.3) / 4
for i, (k, lab) in enumerate(names):
    x = R + i * (pw + 0.1)
    picture(IMG / f"plan_{k}.png", x + 0.06, py, w=pw - 0.12, border=False)
    tb(x, py + (pw - 0.12) * 1.03 + 0.01, pw, 0.22, lab, size=9, bold=k == "optimized",
       color=SIGNAL if k == "optimized" else INK, align=PP_ALIGN.CENTER, inset=0)
caption(R, py + (pw - 0.12) * 1.03 + 0.22, CW, "Fig. 4  Rack visits per month (grey = rare, red = very frequent); dock at the bottom.")

# ---- SDGs
sy = 10.23
header(R - 0.05, sy, "SDGs Addressed:", w=2.4)
for i, (num, col, lab) in enumerate([("8", "A21942", "Decent work: less walking, less fatigue"),
                                      ("9", "FD6925", "Industry & innovation: digital twin + IoT"),
                                      ("12", "BF8B2E", "Responsible production: leaner logistics")]):
    yy = sy + 0.4 + i * 0.27
    box(R + 0.02, yy, 0.42, 0.23, col, text=f"SDG {num}", size=8, kind=MSO_SHAPE.RECTANGLE)
    tb(R + 0.5, yy - 0.01, CW - 0.5, 0.25, lab, size=10.5, align=PP_ALIGN.LEFT, inset=0.02)

# ---- Conclusion
ky = 11.47
header(R - 0.05, ky, "Conclusion/ Summary")
tb(R, ky + 0.36, CW, 1.62, [[
    ("The textbook ABC rule saves only 2.5% because orders average 7 lines spread across the floor. Using ", {}),
    ("what customers buy together", {"bold": True}), (" saves 11.3%, and simulated annealing 15.8% (34 km ≈ 8 picker-hours a month); "
     "the gains hold on unseen orders (13.9%). Smart batching plus the optimized layout cuts walking by 68%. With a ", {}),
    ("team", {"bold": True}), (", narrow aisles lose 23–35% of capacity to queueing, which erodes about half the slotting gain: "
     "aisle width and batching matter as much as slotting. Next: live sensor input and dynamic re-slotting.", {})]], border=True)

# ---- Contact + references (template order)
cy2 = 13.72
line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(R - 0.05), Inches(cy2 - 0.04), Inches(R + CW), Inches(cy2 - 0.04))
line.line.color.rgb = rgb(INK)
line.line.width = Pt(0.75)
tb(R - 0.05, cy2, CW + 0.05, 0.5, [[("Contact Details : ", {"bold": True, "color": NAVY, "size": 17.8}),
                                     ("rohithkrishna.a2025@vitstudent.ac.in", {"size": 12, "color": "0563C1"})],
                                    [("jaidev.a2025@vitstudent.ac.in   ·   jaffino.g@vit.ac.in (guide)", {"size": 12, "color": "0563C1"})]],
   align=PP_ALIGN.LEFT, inset=0.02)
header(R - 0.05, cy2 + 0.55, "Acknowledgments/ References")
refs = [
    "Chen, D. (2019). Online Retail II [Dataset]. UCI Machine Learning Repository. https://doi.org/10.24432/C5CG6D",
    "de Koster, R., Le-Duc, T., & Roodbergen, K. J. (2007). Design and control of warehouse order picking: A literature review. "
    "European Journal of Operational Research, 182(2), 481–501.",
    "Kirkpatrick, S., Gelatt, C. D., & Vecchi, M. P. (1983). Optimization by simulated annealing. Science, 220(4598), 671–680.",
    "Gue, K. R., Meller, R. D., & Skufca, J. D. (2006). The effects of pick density on order picking areas with narrow aisles. "
    "IIE Transactions, 38(10), 859–868.",
]
rb = tb(R, cy2 + 0.9, CW, 1.55, [[(r, {})] for r in refs], size=9, ls=0.9, after=1.5)
for p in rb.text_frame.paragraphs:  # numbered list, as in the template
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "200000")
    pPr.set("indent", "-200000")
    an = etree.SubElement(pPr, qn("a:buAutoNum"))
    an.set("type", "arabicPeriod")

prs.save(OUT)
print("saved", OUT)
